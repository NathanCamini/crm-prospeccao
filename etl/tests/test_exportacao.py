"""Exportação CSV/XLSX/PDF e descrição de filtros."""

from __future__ import annotations

import csv
import io
from datetime import timedelta

import openpyxl
import pytest
from pypdf import PdfReader

from prospeccao.app_store import AppStore, ContatoNovo
from prospeccao.consulta import COLUNAS_PADRAO, Repositorio
from prospeccao.domain.filtros import Filtros, Localidade, Periodo
from prospeccao.reports.descricao import descrever_filtros
from prospeccao.reports.gerar import gerar_arquivo
from tests.conftest import REFERENCIA

A, D = "00012345000165", "44555666000181"


@pytest.fixture()
def ambiente(carga_amostra, tmp_path):
    cfg, _ = carga_amostra
    store = AppStore(tmp_path / "app.sqlite")
    store.adicionar_contato(ContatoNovo(A, "site", "https://alfaobras.com.br", "alfaobras.com.br", "site oficial",
                                        "alta", "https://alfaobras.com.br/contato", data_busca="2026-10-01T10:00:00"))
    store.adicionar_contato(ContatoNovo(A, "telefone", "(53) 3999-1234", "5339991234", "busca manual", "media"))
    store.definir_status(D, "em_contato", "Ligar terça")
    with Repositorio(cfg.gold_db, app=store) as repo:
        yield repo, store


def _gerar(ambiente, formato: str, **kw) -> bytes:
    repo, store = ambiente
    buffer = io.BytesIO()
    gerar_arquivo(formato, buffer, repo, store, kw.pop("filtros", Filtros()), referencia=REFERENCIA, usuario="ana", **kw)
    return buffer.getvalue()


def _csv(dados: bytes, sep: str = ";") -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(dados.decode("utf-8-sig")), delimiter=sep))


def test_csv_utf8_com_bom_e_ponto_e_virgula(ambiente) -> None:
    dados = _gerar(ambiente, "csv", colunas=(*COLUNAS_PADRAO, "capital_social"))
    assert dados.startswith(b"\xef\xbb\xbf")  # BOM: Excel abre com acentuação correta
    linhas = _csv(dados)
    assert len(linhas) == 4
    padaria = next(l for l in linhas if l["Nome fantasia"] == "PADARIA SÃO JOÃO")  # acentos preservados
    assert padaria["CNPJ"] == "44.555.666/0001-81"
    assert padaria["Capital social"] == "5000,00"  # decimal com vírgula
    assert padaria["Status da prospecção"] == "Em contato"
    assert padaria["Observação da prospecção"] == "Ligar terça"


def test_csv_mantem_contato_original_e_separa_complementares(ambiente) -> None:
    linhas = {l["CNPJ"]: l for l in _csv(_gerar(ambiente, "csv"))}
    padaria, alfa = linhas["44.555.666/0001-81"], linhas["00.012.345/0001-65"]
    assert padaria["Telefone 1 (Receita)"] == " 99999000"            # exatamente como na Receita
    assert padaria["E-mail (Receita)"] == "Joao.Silva@Exemplo.com"
    assert alfa["E-mail (Receita)"] == "CONTATO@ALFA-OBRAS.COM.BR"
    assert alfa["Sites complementares"].startswith("https://alfaobras.com.br (confiança alta;")
    assert "01/10/2026" in alfa["Sites complementares"]
    assert alfa["Telefones complementares"].startswith("(53) 3999-1234 (confiança média;")
    assert padaria["Sites complementares"] == ""


def test_csv_sempre_inclui_contato_da_receita(ambiente) -> None:
    cabecalho = _csv(_gerar(ambiente, "csv", colunas=("razao_social",)))[0].keys()
    assert {"CNPJ", "DDD 1 (Receita)", "Telefone 1 (Receita)", "E-mail (Receita)"} <= set(cabecalho)


def test_csv_com_virgula_usa_ponto_decimal(ambiente) -> None:
    linhas = _csv(_gerar(ambiente, "csv", colunas=("capital_social",), separador=","), sep=",")
    assert "5000.00" in {l["Capital social"] for l in linhas}


def test_xlsx_tres_abas(ambiente) -> None:
    wb = openpyxl.load_workbook(io.BytesIO(_gerar(ambiente, "xlsx")))
    assert wb.sheetnames == ["Empresas", "Filtros", "Indicadores"]
    empresas = list(wb["Empresas"].iter_rows(values_only=True))
    assert len(empresas) == 5  # cabeçalho + 4
    cab = empresas[0]
    padaria = next(r for r in empresas[1:] if r[cab.index("CNPJ")] == "44.555.666/0001-81")
    assert padaria[cab.index("Telefone 1 (Receita)")] == " 99999000"
    assert padaria[cab.index("Início da atividade")].date() == REFERENCIA  # data de verdade, não texto
    filtros = dict(r for r in wb["Filtros"].iter_rows(values_only=True) if r[0])
    assert filtros["Base: mês de referência"] == "09/2026"
    assert filtros["Total de registros"] == "4"
    assert "últimos 90 dias" in filtros["Abertura (início de atividade)"]
    assert "descadastro" in filtros["Uso dos dados (LGPD)"]
    indicadores = [r for r in wb["Indicadores"].iter_rows(values_only=True)]
    assert ("4399103", "Obras de alvenaria", 2) in indicadores


def test_pdf_com_filtros_data_da_base_total_e_top(ambiente) -> None:
    dados = _gerar(ambiente, "pdf", top_n=3)
    assert dados.startswith(b"%PDF")
    texto = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(dados)).pages)
    for trecho in ("Filtros aplicados", "Base: mês de referência", "09/2026", "Total de registros",
                   "Top 3 prospects", "Top 10 atividades", "Por porte", "Por mês de abertura", "descadastro",
                   "PADARIA SÃO JOÃO", "Em contato"):
        assert trecho in texto, trecho


def test_somente_cnpjs_na_exportacao(ambiente) -> None:
    assert len(_csv(_gerar(ambiente, "csv", somente_cnpjs=[D]))) == 1


def test_descricao_dos_filtros() -> None:
    f = Filtros(
        localidades=(Localidade("RS", "9701", ("CENTRO",)), Localidade("SC")),
        periodo=Periodo(None, REFERENCIA - timedelta(days=5), REFERENCIA), cnae_prefixos=("43",),
        incluir_cnaes_secundarios=True, capital_min=1000.0, simples=False, status=("em_contato",),
    )
    itens = dict(descrever_filtros(f, REFERENCIA, {"9701": "CIDADE ALFA/RS"}))
    assert itens["Localidade"] == "CIDADE ALFA/RS (bairros: CENTRO); SC (UF inteira)"
    assert itens["Abertura (início de atividade)"] == "04/10/2026 a 09/10/2026"
    assert itens["CNAE"] == "prefixos 43 (inclui secundários)"
    assert itens["Capital social"] == "R$ 1.000,00 a sem máximo"
    assert itens["Simples Nacional"] == "não optante"
    assert itens["Status de prospecção"] == "Em contato"
