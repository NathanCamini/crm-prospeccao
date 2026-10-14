"""Exportação do resultado em CSV (UTF-8 com BOM) e XLSX (resultados, filtros, indicadores).

Inclui sempre o contato original da Receita (sem alteração) e, em colunas separadas,
os contatos complementares e o status de prospecção.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any, BinaryIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from prospeccao.app_store import AppStore
from prospeccao.consulta import Repositorio
from prospeccao.domain.cnpj import formatar_cnpj
from prospeccao.domain.filtros import Filtros
from prospeccao.domain.status import STATUS, STATUS_PADRAO
from prospeccao.web.rotulos import CAMPOS_CONTATO_RECEITA, rotulo

LIMITE_LINHAS_XLSX = 1_048_575
AVISO_LGPD = (
    "Dados públicos do CNPJ (Receita Federal) usados para prospecção comercial B2B. Todo contato deve "
    "oferecer um canal de descadastro e atender pedidos de exclusão. Os contatos têm data de referência "
    "e podem estar desatualizados."
)
CONTATO_RECEITA_EXPORTADO = ("ddd_receita_1", "telefone_receita_1", "ddd_receita_2", "telefone_receita_2", "email_receita")
COLUNAS_EXTRAS = {
    "status_prospeccao": "Status da prospecção",
    "observacao_prospeccao": "Observação da prospecção",
    "compl_telefone": "Telefones complementares",
    "compl_email": "E-mails complementares",
    "compl_site": "Sites complementares",
    "compl_rede_social": "Redes sociais complementares",
}
_CONFIANCA = {"alta": "alta", "media": "média", "baixa": "baixa"}


@dataclass
class Metadados:
    total: int
    filtros: list[tuple[str, str]]
    mes_referencia: str | None
    data_carga: datetime | None
    usuario: str | None = None
    gerado_em: datetime = field(default_factory=lambda: datetime.now().replace(microsecond=0))

    def linhas(self) -> list[tuple[str, str]]:
        ref = self.mes_referencia or "—"
        if self.mes_referencia and "-" in self.mes_referencia:
            ano, mes = self.mes_referencia.split("-", 1)
            ref = f"{mes}/{ano}"
        return [
            ("Base: mês de referência", ref),
            ("Base: carregada em", f"{self.data_carga:%d/%m/%Y %H:%M}" if self.data_carga else "—"),
            ("Gerado em", f"{self.gerado_em:%d/%m/%Y %H:%M}"),
            ("Gerado por", self.usuario or "—"),
            ("Total de registros", f"{self.total:,}".replace(",", ".")),
        ]


def colunas_exportacao(colunas_visiveis: Iterable[str]) -> tuple[str, ...]:
    """Colunas visíveis + CNPJ + contato original da Receita (sempre)."""
    colunas = list(dict.fromkeys(["cnpj", *colunas_visiveis]))
    colunas += [c for c in CONTATO_RECEITA_EXPORTADO if c not in colunas]
    return tuple(colunas)


def _resumo_contato(c: dict[str, Any]) -> str:
    data = (c.get("data_busca") or "")[:10]
    if len(data) == 10:
        data = f"{data[8:10]}/{data[5:7]}/{data[:4]}"
    fonte = c.get("fonte_url") or c.get("fonte_nome") or ""
    return f"{c['valor']} (confiança {_CONFIANCA.get(c['confianca'], c['confianca'])}; {fonte}; {data})"


def acrescentar_extras(linhas: list[dict[str, Any]], store: AppStore | None) -> list[dict[str, Any]]:
    """Status e contatos complementares (em colunas próprias) para um lote de linhas."""
    cnpjs = [l["cnpj"] for l in linhas]
    status = store.status_de(cnpjs) if store else {}
    contatos = store.contatos_por_cnpjs(cnpjs) if store else {}
    for linha in linhas:
        st = status.get(linha["cnpj"])
        linha["status_prospeccao"] = STATUS[st["status"] if st else STATUS_PADRAO]
        linha["observacao_prospeccao"] = (st or {}).get("observacao")
        por_tipo: dict[str, list[str]] = {}
        for c in contatos.get(linha["cnpj"], []):
            por_tipo.setdefault(c["tipo"], []).append(_resumo_contato(c))
        for tipo in ("telefone", "email", "site", "rede_social"):
            linha[f"compl_{tipo}"] = " | ".join(por_tipo.get(tipo, [])) or None
    return linhas


def linhas_exportacao(
    repo: Repositorio,
    store: AppStore | None,
    filtros: Filtros,
    colunas: tuple[str, ...],
    referencia: date | None = None,
    somente_cnpjs: list[str] | None = None,
    limite: int | None = None,
) -> Iterator[dict[str, Any]]:
    for lote in repo.iterar(filtros, colunas, referencia, somente_cnpjs, limite):
        yield from acrescentar_extras(lote, store)


def _cabecalho(coluna: str) -> str:
    return COLUNAS_EXTRAS.get(coluna) or rotulo(coluna)


def _texto_csv(coluna: str, valor: Any, decimal: str) -> str:
    if coluna in CAMPOS_CONTATO_RECEITA:
        return "" if valor is None else str(valor)  # exatamente como veio da Receita
    if valor is None:
        return ""
    if coluna == "cnpj":
        return formatar_cnpj(valor)  # texto com máscara: o Excel não remove os zeros à esquerda
    if coluna == "cep" and isinstance(valor, str) and len(valor) == 8:
        return f"{valor[:5]}-{valor[5:]}"
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, Decimal | float):
        return f"{valor:.2f}".replace(".", decimal)
    if isinstance(valor, list):
        return ", ".join(valor)
    return str(valor)


def escrever_csv(linhas: Iterable[dict[str, Any]], destino: BinaryIO, colunas: tuple[str, ...], separador: str = ";") -> int:
    """CSV UTF-8 com BOM (abre no Excel com acentos). Decimal com vírgula quando o separador é ';'."""
    decimal = "," if separador == ";" else "."
    texto = io.TextIOWrapper(destino, encoding="utf-8-sig", newline="", write_through=True)
    escritor = csv.writer(texto, delimiter=separador, quoting=csv.QUOTE_MINIMAL)
    todas = (*colunas, *COLUNAS_EXTRAS)
    escritor.writerow([_cabecalho(c) for c in todas])
    total = 0
    for linha in linhas:
        escritor.writerow([_texto_csv(c, linha.get(c), decimal) for c in todas])
        total += 1
    texto.flush()
    texto.detach()  # não fecha o destino
    return total


def _valor_xlsx(coluna: str, valor: Any) -> Any:
    if coluna in CAMPOS_CONTATO_RECEITA:
        return valor  # texto, sem alteração
    if valor is None:
        return None
    if coluna == "cnpj":
        return formatar_cnpj(valor)
    if coluna == "cep" and isinstance(valor, str) and len(valor) == 8:
        return f"{valor[:5]}-{valor[5:]}"
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, list):
        return ", ".join(valor)
    return valor


def _linha_titulo(ws: Any, valores: list[str]) -> None:
    celulas = []
    for v in valores:
        c = WriteOnlyCell(ws, value=v)
        c.font = Font(bold=True)
        celulas.append(c)
    ws.append(celulas)


def escrever_xlsx(
    linhas: Iterable[dict[str, Any]],
    destino: BinaryIO,
    colunas: tuple[str, ...],
    meta: Metadados,
    indicadores: dict[str, Any],
) -> int:
    """Três abas: Empresas, Filtros (com dados da base e aviso LGPD) e Indicadores."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Empresas")
    todas = (*colunas, *COLUNAS_EXTRAS)
    _linha_titulo(ws, [_cabecalho(c) for c in todas])
    total, cortado = 0, False
    for linha in linhas:
        if total >= LIMITE_LINHAS_XLSX:
            cortado = True
            break
        celulas = []
        for c in todas:
            cel = WriteOnlyCell(ws, value=_valor_xlsx(c, linha.get(c)))
            if isinstance(cel.value, date):
                cel.number_format = "DD/MM/YYYY"
            elif c == "capital_social" and cel.value is not None:
                cel.number_format = '"R$" #,##0.00'
            celulas.append(cel)
        ws.append(celulas)
        total += 1

    wf = wb.create_sheet("Filtros")
    _linha_titulo(wf, ["Item", "Valor"])
    for item in meta.linhas():
        wf.append(list(item))
    wf.append([])
    _linha_titulo(wf, ["Filtro", "Valor"])
    for item in meta.filtros:
        wf.append(list(item))
    wf.append([])
    if cortado:
        wf.append(["Atenção", f"Resultado cortado em {LIMITE_LINHAS_XLSX:,} linhas (limite do Excel). Use CSV."])
    wf.append(["Uso dos dados (LGPD)", AVISO_LGPD])

    wi = wb.create_sheet("Indicadores")
    _linha_titulo(wi, ["Total de empresas", str(indicadores.get("total", total))])
    wi.append([])
    _linha_titulo(wi, ["Top 10 CNAE", "Descrição", "Empresas"])
    for c in indicadores.get("top_cnae", []):
        wi.append([c["codigo"], c["descricao"], c["total"]])
    wi.append([])
    _linha_titulo(wi, ["Porte", "Empresas"])
    for p in indicadores.get("por_porte", []):
        wi.append([p["porte"], p["total"]])
    wi.append([])
    _linha_titulo(wi, ["Mês de abertura", "Empresas"])
    for m in indicadores.get("por_mes", []):
        wi.append([m["mes"], m["total"]])
    if indicadores.get("por_municipio"):
        wi.append([])
        _linha_titulo(wi, ["Município", "UF", "Empresas"])
        for m in indicadores["por_municipio"]:
            wi.append([m["municipio"], m["uf"], m["total"]])
    wb.save(destino)
    return total
